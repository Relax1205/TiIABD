import copy, os, sys
import docx
from docx.shared import Pt, Cm, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls
from PIL import Image, ImageDraw, ImageFont

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TEMPLATE = os.path.join(ROOT, "Отчёт_ПР2_Система_хранения_данных.docx")
OUT = os.path.join(ROOT, "Отчёт_ПР3_Распределённая_обработка_Spark.docx")
IMG = os.path.join(ROOT, "_build", "img")
os.makedirs(IMG, exist_ok=True)

# ---------------------------------------------------------------- картинки с кодом
FONT = ImageFont.truetype(r"C:\Windows\Fonts\consola.ttf", 22)

def code_image(name, code):
    lines = code.strip("\n").split("\n")
    lh, pad = 31, 22
    w = max(int(FONT.getlength(l)) for l in lines) + 2 * pad
    w = max(w, 900)
    h = lh * len(lines) + 2 * pad
    im = Image.new("RGB", (w, h), "#F6F8FA")
    d = ImageDraw.Draw(im)
    d.rectangle([0, 0, w - 1, h - 1], outline="#C8CDD2", width=2)
    for i, l in enumerate(lines):
        col = "#6A737D" if l.strip().startswith("#") else "#1F2328"
        d.text((pad, pad + i * lh), l, font=FONT, fill=col)
    path = os.path.join(IMG, name + ".png")
    im.save(path)
    return path, w, h

# ---------------------------------------------------------------- документ-шаблон
doc = docx.Document(TEMPLATE)
body = doc.element.body
paras = doc.paragraphs

# титульный лист
for p in paras[:25]:
    t = p.text.strip()
    if t == "Практическая работа № 2":
        p.runs[0].text = "Практическая работа № 3"
        for r in p.runs[1:]:
            r.text = ""
    elif t == "Система хранения данных":
        p.runs[0].text = "Распределённая модель обработки данных"
        for r in p.runs[1:]:
            r.text = ""

intro = next(p for p in paras if p.text.strip() == "ВВЕДЕНИЕ")
sdt = next(c for c in body if c.tag == qn("w:sdt"))
sect = body.find(qn("w:sectPr"))

# удалить всё, начиная с введения
started = False
for el in list(body):
    if el is intro._p:
        started = True
    if started and el is not sect:
        body.remove(el)
toc_anchor = OxmlElement("w:p")
sdt.addprevious(toc_anchor)
body.remove(sdt)

# ---------------------------------------------------------------- хелперы
def _ppr(p, **kw):
    pPr = p._p.get_or_add_pPr()
    for tag, attrs in kw.items():
        pass
    return pPr

def spacing(p, before=None, after=None, line=None):
    pf = p.paragraph_format
    if before is not None: pf.space_before = Pt(before / 20)
    if after is not None: pf.space_after = Pt(after / 20)
    if line: pf.line_spacing = line

def para(text, first=True, justify=True):
    p = doc.add_paragraph(text)
    pf = p.paragraph_format
    pf.line_spacing = 1.5
    if first: pf.first_line_indent = Emu(709 * 635)
    if justify: p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return p

def h1(text, center=False):
    p = doc.add_paragraph(style="Heading 1")
    p.add_run(text)
    pf = p.paragraph_format
    pf.keep_with_next = True
    pf.line_spacing = 1.5
    pf.space_after = Pt(12)
    if center:
        pf.page_break_before = True
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        pf.space_before = Pt(18)
        pf.first_line_indent = Emu(709 * 635)
    HEADINGS.append((1, text))
    return p

def h2(text):
    p = doc.add_paragraph(style="Heading 2")
    p.add_run(text)
    pf = p.paragraph_format
    pf.keep_with_next = True
    pf.line_spacing = 1.5
    pf.space_before = Pt(12)
    pf.space_after = Pt(12)
    pf.first_line_indent = Emu(709 * 635)
    HEADINGS.append((2, text))
    return p

def shade(p, fill):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    pPr.append(shd)

def box(p, color="BF9000"):
    pPr = p._p.get_or_add_pPr()
    b = parse_xml(
        f'<w:pBdr {nsdecls("w")}>' +
        "".join(f'<w:{s} w:val="dashed" w:sz="8" w:space="6" w:color="{color}"/>' for s in ("top", "left", "bottom", "right")) +
        "</w:pBdr>")
    pPr.append(b)

FIG = [0]
TAB = [0]
HEADINGS = []

def caption_fig(text):
    FIG[0] += 1
    p = doc.add_paragraph(f"Рисунок {FIG[0]} – {text}")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(12)
    return FIG[0]

def nextfig(k=1):
    return FIG[0] + k

def code_fig(name, code, caption, width_cm=16.0):
    path, w, h = code_image(name, code)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    width = min(width_cm, w / 900 * 16.0 + 0.01)
    width = max(width, 8)
    p.add_run().add_picture(path, width=Cm(min(width_cm, w / 60.0)))
    return caption_fig(caption)

def placeholder(title, steps, caption):
    """Жёлтый блок: пользователь удаляет его и вставляет свой скриншот."""
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(f"[МЕСТО ДЛЯ СКРИНШОТА] {title}")
    r.bold = True; r.font.size = Pt(12)
    shade(p, "FFF2CC"); box(p)
    lines = steps if isinstance(steps, list) else [steps]
    for i, s in enumerate(lines):
        q = doc.add_paragraph()
        q.paragraph_format.line_spacing = 1.15
        q.paragraph_format.space_after = Pt(0 if i < len(lines) - 1 else 6)
        q.paragraph_format.keep_with_next = i < len(lines) - 1 or True
        rr = q.add_run(s); rr.font.size = Pt(12)
        shade(q, "FFF2CC"); box(q)
    return caption_fig(caption)

def table(caption, header, rows, widths):
    TAB[0] += 1
    cap = doc.add_paragraph(f"Таблица {TAB[0]} – {caption}")
    cap.paragraph_format.keep_with_next = True
    cap.paragraph_format.space_before = Pt(6)
    cap.paragraph_format.space_after = Pt(6)
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    tbl = t._tbl
    tblPr = tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>' +
        "".join(f'<w:{s} w:val="single" w:sz="4" w:space="0" w:color="auto"/>' for s in ("top", "left", "bottom", "right", "insideH", "insideV")) +
        "</w:tblBorders>")
    tblPr.append(borders)
    data = [header] + rows
    for ri, row in enumerate(data):
        for ci, val in enumerate(row):
            cell = t.cell(ri, ci)
            cell.width = Cm(widths[ci])
            pp = cell.paragraphs[0]
            pp.paragraph_format.space_after = Pt(2)
            pp.paragraph_format.space_before = Pt(2)
            run = pp.add_run(str(val))
            run.font.size = Pt(12)
            run.bold = ri == 0
        if ri == 0:
            trPr = t.rows[0]._tr.get_or_add_trPr()
            trPr.append(OxmlElement("w:tblHeader"))
        cs = OxmlElement("w:cantSplit")
        t.rows[ri]._tr.get_or_add_trPr().append(cs)
    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(6)
    return TAB[0]

def nexttab():
    return TAB[0] + 1

# ================================================================= СОДЕРЖИМОЕ
# введение
h1("ВВЕДЕНИЕ", center=True)
para("Apache Spark – вычислительный движок для параллельной обработки больших наборов табличных данных. Он отделяет описание "
     "преобразований от их выполнения: пользователь строит цепочку операций над DataFrame, а Spark лениво формирует "
     "план и исполняет его по разделам данных. В работе Spark запускается в локальном режиме local[*], при котором "
     "программная модель кластера сохраняется, но используются ресурсы одного компьютера.")
para("Цель работы – освоить базовый цикл распределённой обработки данных средствами PySpark: от чтения данных и "
     "задания схемы до агрегаций, соединений, Spark SQL и записи результата в Parquet.")
para("Задачи работы:")
para("а) подготовить окружение и запустить локальную SparkSession;")
para("б) загрузить таблицы базы ecommerce.sqlite в Spark DataFrame и выполнить учебную последовательность преобразований;")
para("в) выполнить самостоятельный анализ собственного набора данных (продажи торговой компании);")
para("г) сохранить итоговые витрины в Parquet и CSV и сформулировать выводы.")
para("Отчёт оформлен по ГОСТ 7.32–2017 [4].")

# ---- 1
h1("1 Подготовка окружения")
para("Работа выполнялась на Windows 11 с Python 3.12, Java 8 и PySpark 3.5.3 (Java 8 поддерживается веткой Spark 3.5; "
     "Spark 4 требует Java 17 и выше). Код оформлен в двух Jupyter Notebook: учебный по базе ecommerce.sqlite и самостоятельный по "
     "собственным данным. Оба ноутбука выполняются последовательно командой Run All.")
h2("1.1 Установка и проверка версий")
n1 = nextfig(); n2 = nextfig(2)
para(f"Пакеты установлены командами pip, версии Python, Java и PySpark зафиксированы в первой ячейке ноутбука "
     f"(рисунок {n1}). Результат проверки приведён на рисунке {n2}.")
code_fig("c1", '''python -m pip install pyspark==3.5.3 pandas jupyter pyarrow

# первая ячейка ноутбука: проверка версий
import sys, subprocess, pyspark
print("Python :", sys.version.split()[0])
print("PySpark:", pyspark.__version__)
print(subprocess.run(["java", "-version"],
      capture_output=True, text=True).stderr.splitlines()[0])''',
         "Установка пакетов и проверка версий")
placeholder("Вывод версий Python, PySpark и Java",
            ["Что сделать: откройте ноутбук «Практика3_ecommerce.ipynb» в Jupyter/VS Code и найдите ячейку «3.1. Окружение» (первая ячейка с кодом).",
             "Выполните её и сделайте скриншот ячейки вместе с выводом. На снимке должны быть три строки: «Python : 3.12.10», «PySpark: 3.5.3» и «java version 1.8.0_503».",
             "Как вставить: удалите этот жёлтый блок, затем Вставка → Рисунок → ваш файл; выровняйте по центру, ширина не более 16 см. Подпись ниже оставьте."],
            "Версии Python, PySpark и Java")

h2("1.2 Выгрузка таблиц SQLite")
n1 = nextfig(); n2 = nextfig(2); t1 = nexttab()
para(f"Spark не читает файл SQLite напрямую, поэтому использован staging-этап: модуль sqlite3 и pandas выгружают "
     f"7 нужных таблиц в CSV-файлы каталога spark_input (рисунок {n1}). Числа строк, которые должна напечатать программа, "
     f"приведены в таблице {t1}; ожидаемый вывод показан на рисунке {n2}.")
code_fig("c2", '''DB_PATH = Path("ecommerce.sqlite")
STAGING = Path("spark_input")
STAGING.mkdir(exist_ok=True)

TABLES = ["orders", "order_items", "products", "customers",
          "order_payments", "order_reviews",
          "product_category_name_translation"]

with sqlite3.connect(DB_PATH) as conn:
    for table in TABLES:
        part = pd.read_sql_query(f'SELECT * FROM "{table}"', conn)
        part.to_csv(STAGING / f"{table}.csv", index=False, encoding="utf-8")
        print(f"{table:36s} {len(part):>8} rows")''', "Код выгрузки таблиц в staging-каталог")
placeholder("Вывод количества строк по таблицам",
            ["Что сделать: в ноутбуке «Практика3_ecommerce.ipynb» выполните ячейку «3.2. Извлечение таблиц SQLite в staging-каталог» и сделайте скриншот её вывода.",
             "На снимке должны быть 7 строк: orders 99441, order_items 112650, products 32951, customers 99441, order_payments 103886, order_reviews 97621, product_category_name_translation 71.",
             "Как вставить: удалите блок, вставьте рисунок по центру (Вставка → Рисунок), подпись оставьте."],
            "Результат выгрузки таблиц")
table("Таблицы ecommerce, используемые в работе",
      ["Таблица", "Строк", "Назначение"],
      [["orders", "99 441", "Заказы, статусы, даты"],
       ["order_items", "112 650", "Позиции заказов, цена, доставка"],
       ["products", "32 951", "Товары и их категории"],
       ["customers", "99 441", "Клиенты и штаты"],
       ["order_payments", "103 886", "Платежи"],
       ["order_reviews", "97 621", "Отзывы"],
       ["product_category_name_translation", "71", "Перевод названий категорий"]],
      [6.3, 2.6, 6.6])

# ---- 2
h1("2 Учебная часть: анализ базы ecommerce")
h2("2.1 Запуск SparkSession")
n1 = nextfig(); n2 = nextfig(2)
para(f"SparkSession создана в режиме local[*] с 2 ГБ памяти драйвера и 8 разделами после shuffle (рисунок {n1}). "
     f"В этом режиме Driver и Executor работают внутри одного процесса JVM, но модель Spark (разделы, задачи, ленивый план) "
     f"остаётся такой же, как в кластере. Параметры сессии показаны на рисунке {n2}.")
code_fig("c3", '''spark = (
    SparkSession.builder
    .appName("Practice4_Ecommerce")
    .master("local[*]")
    .config("spark.driver.memory", "2g")
    .config("spark.sql.shuffle.partitions", "8")
    .getOrCreate()
)
spark.sparkContext.setLogLevel("WARN")''', "Создание SparkSession")
placeholder("Параметры запущенной SparkSession",
            ["Что сделать: выполните в ноутбуке ячейку «4. Запуск локального Spark» и сделайте скриншот вывода.",
             "На снимке должны быть строки «Master: local[*]», «App: Practice4_Ecommerce», «Spark: 3.5.3».",
             "Как вставить: удалите блок, вставьте рисунок по центру, подпись оставьте."],
            "Параметры SparkSession")

h2("2.2 Чтение данных и схема")
n1 = nextfig(); n2 = nextfig(2)
para(f"Таблица orders прочитана без вывода типов, поэтому все столбцы имеют тип string. Для order_items задана явная схема "
     f"StructType (рисунок {n1}): числовые поля получили типы integer и double. Даты orders затем преобразованы в timestamp "
     f"функцией to_timestamp. Полученные схемы приведены на рисунке {n2}.")
code_fig("c4", '''order_items_schema = StructType([
    StructField("order_id", StringType(), True),
    StructField("order_item_id", IntegerType(), True),
    StructField("product_id", StringType(), True),
    StructField("seller_id", StringType(), True),
    StructField("shipping_limit_date", StringType(), True),
    StructField("price", DoubleType(), True),
    StructField("freight_value", DoubleType(), True),
])
items = (spark.read.option("header", True)
         .schema(order_items_schema)
         .csv(str(STAGING / "order_items.csv")))

for c in timestamp_cols:
    orders = orders.withColumn(
        c, F.to_timestamp(F.col(c), "yyyy-MM-dd HH:mm:ss"))''', "Явная схема order_items и преобразование дат")
placeholder("Схемы orders и order_items",
            ["Что сделать: в ноутбуке выполните ячейки «5.1», «5.2» и «5.3» (чтение orders, схема order_items, преобразование дат).",
             "Сделайте один скриншот (или два, если не помещается) с выводом printSchema: у order_items должны быть типы integer/double, у orders после 5.3 – пять полей timestamp; видны также строки «Строк: 99441» и «Строк: 112650».",
             "Как вставить: удалите блок, вставьте рисунок(ки) по центру, подпись оставьте."],
            "Схемы DataFrame после чтения")

h2("2.3 Преобразования, пропуски и агрегации")
n1 = nextfig(); n2 = nextfig(2); t2 = nexttab()
para(f"Для доставленных заказов вычислены признаки delivery_days и is_late (рисунок {n1}). Проверка пропусков показала, что "
     f"поле order_delivered_customer_date пусто у 2 965 заказов, причём почти все они имеют статусы shipped, canceled, "
     f"unavailable и другие незавершённые: пропуск отражает состояние процесса, поэтому строки не удалялись, "
     f"а для категории товара использована метка unknown. Ключевые результаты приведены в таблице {t2}, "
     f"вывод программы – на рисунке {n2}.")
code_fig("c5", '''delivered = (orders.filter(F.col("order_status") == "delivered")
    .withColumn("delivery_days", F.datediff(
        "order_delivered_customer_date", "order_purchase_timestamp"))
    .withColumn("is_late", F.col("order_delivered_customer_date")
                         > F.col("order_estimated_delivery_date")))

null_counts = orders.select([
    F.sum(F.col(c).isNull().cast("int")).alias(c) for c in timestamp_cols])

products_clean = products.fillna({"product_category_name": "unknown"})

payment_stats = (payments.groupBy("payment_type").agg(
        F.count("*").alias("payments"),
        F.round(F.sum("payment_value"), 2).alias("total_value"),
        F.round(F.avg("payment_value"), 2).alias("avg_value"))
    .orderBy(F.desc("total_value")))''', "Признаки доставки, подсчёт пропусков и агрегация платежей")
placeholder("Результаты преобразований, пропусков и агрегаций",
            ["Что сделать: в ноутбуке выполните ячейки «6.2», «7.1», «7.2» и «8.2» и сделайте скриншот их выводов (можно 2 снимка).",
             "Должно быть видно: «Доставленных заказов: 96478», «Позже ожидаемой даты: 7826»; таблица пропусков (order_delivered_customer_date = 2965); таблица статусов с пропущенной датой доставки (shipped 1107, canceled 619, unavailable 609...); таблица по типам платежа (credit_card 76795 платежей).",
             "Как вставить: удалите блок, вставьте рисунок(ки) по центру, подпись оставьте."],
            "Результаты преобразований, пропусков и агрегаций")
table("Ключевые результаты учебной части",
      ["Показатель", "Значение"],
      [["Заказов delivered / с датой доставки", "96 478 / 96 470"],
       ["Доставлено позже ожидаемой даты", "7 826 (≈ 8,1 %)"],
       ["Пропусков order_delivered_customer_date", "2 965"],
       ["Товаров без категории", "610"],
       ["Позиций с price ≥ 500", "3 239"],
       ["Платежей credit_card", "76 795 (12,54 млн)"]],
      [9.5, 6.0])

h2("2.4 Соединения, Spark SQL и запись результата")
n1 = nextfig(); n2 = nextfig(2)
para(f"Позиции заказов соединены с товарами и переводом категорий двумя LEFT JOIN, число строк до и после соединения осталось "
     f"112 650, а для 1 603 позиций категория осталась unknown. Категорийная витрина построена через Spark SQL и "
     f"повторена средствами DataFrame API – результаты совпали. Витрина из 74 категорий записана в Parquet и CSV (рисунок {n1}) "
     f"и прочитана обратно (рисунок {n2}).")
code_fig("c6", '''sales.createOrReplaceTempView("sales")
category_sql = spark.sql("""
SELECT category, COUNT(*) AS items,
       COUNT(DISTINCT order_id) AS orders,
       ROUND(SUM(price), 2) AS revenue,
       ROUND(AVG(price), 2) AS avg_price
FROM sales GROUP BY category
ORDER BY revenue DESC LIMIT 10
""")

category_sales.write.mode("overwrite").parquet("spark_output/category_sales_parquet")
category_sales.coalesce(1).write.mode("overwrite") \\
    .option("header", True).csv("spark_output/category_sales_csv")
print(spark.read.parquet("spark_output/category_sales_parquet").count())''',
         "Spark SQL и запись результата в Parquet и CSV")
placeholder("Топ категорий и результат записи",
            ["Что сделать: в ноутбуке выполните ячейки «9.1» (JOIN) и «10.1» (SQL), «11. Запись результата» и сделайте скриншоты выводов.",
             "Должно быть видно: «Строк до JOIN: 112650», «Строк после JOIN: 112650», «Позиций категории unknown: 1603»; таблица топ-10 категорий (health_beauty 1258681.34, watches_gifts 1205005.68, bed_bath_table 1036988.68); «Строк после round-trip: 74».",
             "Дополнительно откройте в проводнике папку spark_output и сделайте скриншот: внутри category_sales_parquet и category_sales_csv видны файлы part-00000... и _SUCCESS.",
             "Как вставить: удалите блок, вставьте рисунки по центру, подпись оставьте."],
            "Результаты соединений, SQL и записи")

# ---- 3
h1("3 Самостоятельная работа")
h2("3.1 Набор данных")
t3 = nexttab()
para(f"Для самостоятельной части выбран набор продаж торговой компании (каталог archive): факт «Продажи» и справочники "
     f"«Продукты» и «Клиенты» (таблица {t3}). Данные обезличены: названия организаций и ИНН вымышленные, физические лица "
     f"представлены единственной записью «Частное лицо». Файлы – CSV с разделителем табуляции, десятичной запятой, "
     f"датами формата dd.MM.yyyy и маркером NULL для пустых значений, поэтому staging не нужен: Spark читает CSV напрямую.")
table("Файлы собственного набора данных",
      ["Файл", "Строк", "Содержание"],
      [["Продажи.csv", "865 222", "Факт: дата, филиал, товар, клиент, количество, суммы, себестоимость"],
       ["Продукты.csv", "21 163", "Справочник товаров, брендов и групп"],
       ["Клиенты.csv", "3 838", "Справочник клиентов (3 227 уникальных)"]],
      [3.6, 2.4, 9.5])

h2("3.2 Схема и вычисляемые признаки")
n1 = nextfig(); n2 = nextfig(2)
para(f"Задана явная схема с латинскими именами полей; маркер NULL превращён в null через nullValue. inferSchema с nullValue "
     f"распознал числом только «Количество», а суммы и себестоимость с десятичной запятой остались строками, поэтому типы приводились явно. "
     f"Добавлены признаки net_amount (сумма после скидки), margin, month и is_delivery (рисунок {n1}). "
     f"Схема итогового DataFrame показана на рисунке {n2}.")
code_fig("c7", '''def to_num(col):
    return F.regexp_replace(F.trim(F.col(col)), ",", ".").cast("double")

sales_t = (sales
    .withColumn("purchase_date", F.to_date("purchase_date", "dd.MM.yyyy"))
    .withColumn("month", F.date_format("purchase_date", "yyyy-MM"))
    .withColumn("cost_price", to_num("cost_price"))
    .withColumn("quantity", to_num("quantity"))
    .withColumn("amount", to_num("amount"))
    .withColumn("discount", to_num("discount"))
    .withColumn("net_amount", F.col("amount") - F.col("discount"))
    .withColumn("cost_total", F.col("cost_price") * F.col("quantity"))
    .withColumn("margin", F.col("net_amount") - F.col("cost_total"))
    .withColumn("is_delivery", F.col("address").isNotNull()))''',
         "Приведение типов и вычисляемые признаки")
placeholder("Схема после преобразований",
            ["Что сделать: откройте ноутбук «Практика3_самостоятельная.ipynb», выполните ячейки раздела «3. Чтение основной таблицы и схема» и «4. Transformations».",
             "Сделайте скриншот вывода, где видны: «Строк: 865222», сравнение схем inferSchema и явной схемы и итоговая схема после преобразований (purchase_date: date, amount: double, net_amount: double, is_delivery: boolean).",
             "Как вставить: удалите блок, вставьте рисунок по центру, подпись оставьте."],
            "Схемы собственного набора данных")

h2("3.3 Пропуски")
t4 = nexttab()
para(f"Пропуски найдены в четырёх полях. Решения по каждому из них и их обоснование приведены в таблице {t4}. "
     f"Строки с null не удалялись: пропуски объясняются процессом продаж.")
table("Пропуски и принятые решения",
      ["Поле", "Пропусков", "Причина и решение"],
      [["address", "832 181", "Самовывоз, доставки нет; строки оставлены, создан признак is_delivery"],
       ["cost_price", "86 431", "Себестоимость неизвестна; нулём не заполнялась, маржа считается по известным"],
       ["receipt_no", "3 934", "Отгрузки без чека; заполнено меткой NO_RECEIPT"],
       ["quantity", "100", "Количество не указано; не заполнялось, сумма продажи известна"]],
      [3.0, 2.6, 9.9])

h2("3.4 Агрегации, соединения и Spark SQL")
n1 = nextfig(); n2 = nextfig(2); t5 = nexttab()
para(f"Показатели по филиалам вычислены семью агрегатами. Справочник клиентов содержал 566 дублирующихся ключей, "
     f"поэтому перед соединением он был дедуплицирован: иначе JOIN размножил бы факты. После LEFT JOIN с товарами и клиентами "
     f"и INNER JOIN с клиентами число строк осталось 865 222. Построено три запроса Spark SQL, один из них "
     f"с GROUP BY и сортировкой по агрегату (рисунок {n1}). Результаты по филиалам – в таблице {t5}, "
     f"вывод запросов – на рисунке {n2}.")
code_fig("c8", '''clients = clients_raw.groupBy("client_id") \\
    .agg(F.min("client_name").alias("client_name"))

sales_p = (sales_c.alias("s")
    .join(products.alias("p"),
          F.col("s.product_key") == F.col("p.product_key"), "left")
    .select("s.*", "p.product_name", "p.brand", "p.product_group"))

sales_full = (sales_p.alias("s")
    .join(clients.alias("c"),
          F.col("s.client_id") == F.col("c.client_id"), "left")
    .select("s.*", "c.client_name"))

sales_full.createOrReplaceTempView("sales_full")
top_groups = spark.sql("""
SELECT group_l1, COUNT(*) AS lines,
       ROUND(SUM(net_amount), 0) AS net_revenue
FROM sales_full GROUP BY group_l1
ORDER BY net_revenue DESC LIMIT 10""")''', "JOIN со справочниками и запрос Spark SQL")
placeholder("Результаты агрегаций, JOIN и SQL",
            ["Что сделать: в ноутбуке «Практика3_самостоятельная.ipynb» выполните ячейки «6.1. Показатели по филиалам», «7.2. Клиенты (INNER JOIN)» и первый запрос из раздела «8. Spark SQL».",
             "Сделайте скриншот (при необходимости два): таблица по филиалам (ТЗ-2 net_revenue ≈ 5.81E7, ТЗ-1 ≈ 4.32E7, ТЗ-3 ≈ 2.58E7); строки «до JOIN: 865222 / INNER JOIN: 865222 / LEFT JOIN: 865222»; топ-10 групп (Живопись ≈ 2.54E7, Графические материалы ≈ 1.40E7).",
             "Как вставить: удалите блок, вставьте рисунок(ки) по центру, подпись оставьте."],
            "Результаты агрегаций, соединений и запросов")
table("Показатели по филиалам (чистая выручка после скидок)",
      ["Филиал", "Строк", "Выручка, млн", "Средняя скидка, %"],
      [["ТЗ-2 (Ленинградское шоссе)", "431 372", "58,12", "1,02"],
       ["ТЗ-1 (ул. Фадеева)", "238 046", "43,23", "1,35"],
       ["ТЗ-3 (Курский)", "190 730", "25,75", "1,33"],
       ["Остальные (3 подразделения)", "5 074", "2,37", "0–3,61"]],
      [6.2, 2.6, 3.4, 3.3])

h2("3.5 Экспорт результата")
n2 = nextfig(1)
para(f"Итоговая витрина «месяц × филиал × группа товаров» (1 780 строк) записана в Parquet, агрегат по филиалам – "
     f"дополнительно в CSV. После повторного чтения количество строк (1 780), имена и типы столбцов и сумма net_revenue "
     f"(129,46 млн) совпали с исходными (рисунок {n2}).")
placeholder("Round-trip витрины",
            ["Что сделать: в ноутбуке «Практика3_самостоятельная.ipynb» выполните ячейки раздела «9. Итоговая витрина и экспорт» (запись и «Повторное чтение Parquet»).",
             "Сделайте скриншот вывода: «Строк в витрине: 1780», «Строк до записи: 1780», «Строк после чтения: 1780», «Имена и типы столбцов совпадают: True», «равны: True».",
             "Как вставить: удалите блок, вставьте рисунок по центру, подпись оставьте."],
            "Результат повторного чтения Parquet")

# заключение
h1("ЗАКЛЮЧЕНИЕ", center=True)
para("В ходе работы освоен базовый цикл распределённой обработки данных в Apache Spark, все задачи выполнены. "
     "Запущена SparkSession в режиме local[*], таблицы ecommerce.sqlite прочитаны в DataFrame с явной и неявной схемами, "
     "выполнены преобразования, обработка пропусков, агрегации, соединения и запросы Spark SQL.")
para("В учебной части показано, что пропуск даты доставки связан со статусом заказа, поэтому удалять такие строки нельзя, "
     "а около 8,1 % доставленных заказов пришли позже ожидаемой даты. LEFT JOIN сохранил все 112 650 позиций заказов.")
para("В самостоятельной части проанализированы продажи торговой компании (865 тыс. строк): почти 96 % строк и около 102 млн из 129 млн "
     "выручки приходится на розничного покупателя, основную выручку дают три торговых зала, наиболее крупная товарная группа – «Живопись». "
     "Ограничения анализа: расчёт маржи ориентировочный (у 13 % строк с известной себестоимостью маржа отрицательна, "
     "а у 10 % строк себестоимость неизвестна), а данные охватывают неполный период с июня 2020 года.")

h1("СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ", center=True)
for s in [
    "1. Apache Spark Documentation / The Apache Software Foundation. – URL: https://spark.apache.org/docs/3.5.3/ (дата обращения: 19.09.2026).",
    "2. PySpark API Reference / The Apache Software Foundation. – URL: https://spark.apache.org/docs/3.5.3/api/python/ (дата обращения: 19.09.2026).",
    "3. Практическая работа № 3. Распределённая модель обработки данных: методические материалы. – РТУ МИРЭА, 2026.",
    "4. ГОСТ 7.32–2017. Система стандартов по информации, библиотечному и издательскому делу. Отчет о научно-исследовательской работе. Структура и правила оформления. – М.: Стандартинформ, 2017.",
    "5. pandas Documentation / pandas via NumFOCUS. – URL: https://pandas.pydata.org/docs/ (дата обращения: 19.09.2026).",
]:
    p = doc.add_paragraph(s)
    p.paragraph_format.line_spacing = 1.5
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

# ---------------------------------------------------------------- оглавление (поле TOC)
def run_fld(kind):
    r = OxmlElement("w:r")
    f = OxmlElement("w:fldChar"); f.set(qn("w:fldCharType"), kind); r.append(f)
    return r

def run_text(text, instr=False):
    r = OxmlElement("w:r")
    t = OxmlElement("w:instrText" if instr else "w:t")
    t.set(qn("xml:space"), "preserve"); t.text = text; r.append(t)
    return r

prev = toc_anchor
for i, (lvl, text) in enumerate(HEADINGS):
    p = OxmlElement("w:p")
    pPr = OxmlElement("w:pPr")
    ind = OxmlElement("w:ind"); ind.set(qn("w:left"), "0" if lvl == 1 else "440"); pPr.append(ind)
    sp = OxmlElement("w:spacing"); sp.set(qn("w:after"), "60"); pPr.append(sp)
    p.append(pPr)
    if i == 0:
        p.append(run_fld("begin")); p.append(run_text(' TOC \\o "1-2" \\h \\z \\u ', True)); p.append(run_fld("separate"))
    p.append(run_text(text))
    if i == len(HEADINGS) - 1:
        p.append(run_fld("end"))
    prev.addnext(p); prev = p

# обновление полей при открытии
settings = doc.settings.element
uf = OxmlElement("w:updateFields"); uf.set(qn("w:val"), "true")
settings.insert(0, uf)

# убрать неиспользуемые картинки шаблона
xml = body.xml if hasattr(body, "xml") else ""
used = set()
for el in body.iter():
    for k, v in el.attrib.items():
        if k.endswith("}embed") or k.endswith("}id") or k.endswith("}link"):
            used.add(v)
part = doc.part
for rid, rel in list(part.rels.items()):
    if rel.reltype.endswith("/image") and rid not in used:
        part.drop_rel(rid) if hasattr(part, "drop_rel") else None
        part.rels.pop(rid, None)

doc.save(OUT)
print("saved", OUT, "figures:", FIG[0], "tables:", TAB[0])
