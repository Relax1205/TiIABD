import json
from docx import Document
from docx.shared import Pt

R = json.load(open("results.json", encoding="utf-8"))
S = R["self"]
d = Document()
d.styles["Normal"].font.name = "Times New Roman"
d.styles["Normal"].font.size = Pt(12)

def f(x, n=3): return f"{x:.{n}f}".replace(".", ",")
def fi(x): return f"{x:,}".replace(",", " ")
def table(title, head, rows):
    d.add_paragraph(title).runs[0].bold = True
    t = d.add_table(rows=1, cols=len(head)); t.style = "Table Grid"
    for i, h in enumerate(head): t.rows[0].cells[i].text = h
    for r in rows:
        c = t.add_row().cells
        for i, v in enumerate(r): c[i].text = str(v)
    d.add_paragraph()
def p(text): d.add_paragraph(text)
def h(text, l=1): d.add_heading(text, l)

d.add_heading("Отчёт по практической работе №4", 0)
p("Тема: исследование выполнения Spark-задач (Jobs/Stages/Tasks, partitions, shuffle, JOIN-стратегии, cache/persist, small files). "
  "Все числа получены запуском Jupyter Notebook «Практика_4.ipynb» (Run All, 0 ошибок); сырые данные — results.json.")

h("1. Стенд")
table("Таблица 2. Паспорт экспериментального стенда", ["Параметр", "Фактическое значение"], list(R["stand"].items()))

h("2. Данные")
p("Использованы таблицы geolocation (1 000 163), order_items (112 650), products (32 951), sellers (3 095), "
  "product_category_name_translation (71); для самостоятельной части — orders (99 441) и order_payments (103 886). "
  f"Исходное число партиций geolocation при чтении Parquet: {R['input_partitions']}. Первая стадия группировки по партициям выполнялась соответствующим числом Tasks.")

h("3. Ленивые вычисления")
L = R["lazy"]
p(f"Цепочка select→filter→withColumn построена за {f(L['build_s'],4)} с, число Jobs в группе P4_LAZY после неё — {L['jobs_after_transformations']}. "
  f"После count() появился {L['jobs_after_action']} Job, результат {fi(L['rows'])} строк за {f(L['count_s'])} с. "
  "Замер вокруг transformations измеряет только построение плана в драйвере.")
table("Таблица 3. Пять наиболее представленных штатов в geolocation", ["state", "Строк", "Доля, %"],
      [(s, fi(n), f(pc, 2)) for s, n, pc in R["top_states"]])

h("4. repartition() и coalesce()")
rows = [(k, fi(v["min"]), fi(v["max"]), f(v["mean"], 1), str(v["sizes"])) for k, v in R["profile"].items()]
table("Таблица 4. Распределение partitions (skewed16 → 4)", ["Вариант", "min rows", "max rows", "mean rows", "Размеры партиций"], rows)
rp, co = R["profile"]["repartition(4)"], R["profile"]["coalesce(4)"]
p(f"repartition(4) добавляет Exchange RoundRobinPartitioning(4) и выравнивает партиции (разброс {fi(rp['min'])}–{fi(rp['max'])}). "
  f"coalesce(4) использует оператор Coalesce без полного shuffle и сохраняет перекос: {fi(co['min'])}–{fi(co['max'])} строк, "
  f"т.к. объединяет соседние партиции, а исходная раскладка по geolocation_state была перекошена (SP = {f(R['top_states'][0][2],2)} %).")

h("5. groupBy и shuffle")
table("Таблица 5. Пять продавцов с максимальной суммой price", ["seller_id", "items", "revenue", "avg_price"],
      [(a, fi(b), f(c, 2), f(e, 2)) for a, b, c, e in R["top_sellers"]])
G = R["ui_groupby"]
p(f"Job(ов): {G['Jobs']}, Stages: {G['Stages']}, Tasks по стадиям: {G['Tasks по Stages']}. "
  f"Shuffle Write = {f(G['Shuffle Write, KiB'],1)} КиБ, Shuffle Read = {f(G['Shuffle Read, KiB'],1)} КиБ. "
  "В плане: HashAggregate (partial) → Exchange hashpartitioning(seller_id, 8) → HashAggregate → Sort.")

h("6. Sort Merge Join и Broadcast Hash Join")
a, b = R["join"]["SMJ"], R["join"]["BHJ"]
table("Таблица 6. Сравнение JOIN-экспериментов (order_items ⋈ products)", ["Показатель", "Sort Merge Join", "Broadcast Hash Join"], [
    ("Строк результата", fi(a["rows"]), fi(b["rows"])),
    ("Exchange hashpartitioning в плане", a["exchange"], b["exchange"]),
    ("Sort в плане", a["sorts"], "0"),
    ("Shuffle Write / Read, КиБ", f"{f(a['Shuffle Write, KiB'],1)} / {f(a['Shuffle Read, KiB'],1)}", f"{f(b['Shuffle Write, KiB'],1)} / {f(b['Shuffle Read, KiB'],1)}"),
    ("Jobs / Stages", f"{a['Jobs']} / {a['Stages']}", f"{b['Jobs']} / {b['Stages']}"),
    ("Tasks по Stages", a["Tasks по Stages"], b["Tasks по Stages"]),
    ("Время, с", f(a["time"]), f(b["time"]))])

h("7. cache() / persist()")
nc, c = R["nocache"], R["cache"]
table("Таблица 7. Повторные вычисления с кэшем и без него", ["Вариант", "Материализация, с", "Action 1, с", "Action 2, с"], [
    ("Без persist", "—", f(nc["a1"]), f(nc["a2"])),
    ("MEMORY_AND_DISK", f(c["mat"]), f(c["a1"]), f(c["a2"]))])
st = c["storage"][0]
p(f"В Storage появился объект: {st['partitions']} партиции, {f(st['memory_kib'],1)} КиБ в памяти, {f(st['disk_kib'],1)} КиБ на диске. "
  f"Контрольные суммы кэшированного и некэшированного вариантов совпали: {c['checksum_equal']}. После unpersist(blocking=True) кэш очищен (is_cached=False).")

h("8. Число партиций и время")
table("Таблица 8. Влияние числа входных partitions (N = 2 000 000)", ["p", "Run 1, с", "Run 2, с", "Run 3, с", "Median, с", "S(p)=T(1)/T(p)"],
      [(x["p"], *[f(t) for t in x["runs"]], f(x["median"]), f(x["speedup"], 2)) for x in R["bench"]])
table("Дополнительно: время при разных spark.sql.shuffle.partitions (вход = 8 партиций, 27 групп)", ["q", "Групп", "Время, с"],
      [(x["q"], x["groups"], f(x["time"])) for x in R["shuffle_q"]])

h("9. Партиции и маленькие файлы")
table("Таблица 9. Связь partitions и part-файлов (order_items)", ["p", "part-файлов", "Общий размер, MiB", "Средний файл, KiB"],
      [(x["p"], x["files"], f(x["total_mib"], 2), f(x["avg_kib"], 1)) for x in R["small"]])

h("10. AQE")
p(f"При включённом AQE и shuffle.partitions=32 план содержит AdaptiveSparkPlan: {R['aqe']['adaptive_plan']}; Tasks по стадиям: {R['aqe']['Tasks по Stages']}.")

h("11. Журнал наблюдений Spark UI (данные REST API Spark UI)")
G2 = lambda n, g: (n, g["Jobs"], g["Stages"], g["Tasks в ключевой Stage"], f(g["Shuffle Write, KiB"], 1) + " КиБ", f(g["Shuffle Read, KiB"], 1) + " КиБ")
table("Таблица 10. Журнал наблюдений по Spark UI", ["Эксперимент", "Jobs", "Stages", "Tasks в ключевой Stage", "Shuffle Write", "Shuffle Read"], [
    G2("groupBy seller_id", R["ui_groupby"]), G2("repartition(4)", rp), G2("coalesce(4)", co),
    G2("Sort Merge Join", a), G2("Broadcast Hash Join", b), G2("persist + повторный Action", c)])
p("Скриншоты вкладок Jobs/Stages/SQL/Storage необходимо сделать вручную при запуске ноутбука (Spark UI живёт только до spark.stop()); "
  "числовые показатели этих вкладок получены автоматически через REST API и приведены выше.")

h("12. Самостоятельная работа (вариант C: orders ⋈ order_payments, ключ order_id, группировка по payment_type)")
p(f"Строк: orders {fi(S['orders_rows'])} (партиций {S['orders_parts']}), order_payments {fi(S['payments_rows'])} (партиций {S['payments_parts']}). "
  f"Lazy: Jobs до Action {S['lazy']['jobs_before']}, после count() {S['lazy']['jobs_after']}; строк после filter(payment_value>0): {fi(S['lazy']['rows'])}.")
table("Профиль партиций (payment_type → 16 → 4)", ["Вариант", "min", "max", "mean", "Shuffle Write, КиБ"],
      [(k, fi(v["min"]), fi(v["max"]), f(v["mean"], 1), f(v["Shuffle Write, KiB"], 1)) for k, v in S["prof"].items()])
p(f"Physical plan: Exchange RoundRobinPartitioning(4) у repartition — {S['plans']['repartition(4)']}; Coalesce у coalesce — {S['plans']['coalesce(4)']}.")
table("groupBy payment_type", ["payment_type", "n", "sum_value"], [(x, fi(n), f(v, 2)) for x, n, v in S["groupby_rows"]])
p(f"Exchange hashpartitioning(payment_type, …) в плане: {S['groupby_exchange']}; Shuffle Write/Read: "
  f"{f(S['groupby_ui']['Shuffle Write, KiB'],1)} / {f(S['groupby_ui']['Shuffle Read, KiB'],1)} КиБ.")
sa, sb = S["join"]["SMJ"], S["join"]["BHJ"]
table("SMJ vs BHJ (orders ⋈ order_payments)", ["Показатель", "Sort Merge Join", "Broadcast Hash Join"], [
    ("Строк результата", fi(sa["rows"]), fi(sb["rows"])), ("Exchange hashpartitioning", sa["exchange"], sb["exchange"]),
    ("Jobs / Stages", f"{sa['Jobs']} / {sa['Stages']}", f"{sb['Jobs']} / {sb['Stages']}"),
    ("Shuffle Write, КиБ", f(sa["Shuffle Write, KiB"], 1), f(sb["Shuffle Write, KiB"], 1)), ("Время, с", f(sa["time"]), f(sb["time"]))])
sc_ = S["cache"]
table("persist на самостоятельном DataFrame", ["Вариант", "Материализация, с", "Action 1, с", "Action 2, с"],
      [("Без persist", "—", f(sc_["nc1"]), f(sc_["nc2"])), ("MEMORY_AND_DISK", f(sc_["mat"]), f(sc_["c1"]), f(sc_["c2"]))])
table("Серия partitions (3 повтора, медиана)", ["p", "Run 1", "Run 2", "Run 3", "Median, с", "S(p)"],
      [(x["p"], *[f(t) for t in x["runs"]], f(x["median"]), f(x["speedup"], 2)) for x in S["bench"]])
table("Запись order_payments", ["p", "part-файлов", "Общий размер, MiB", "Средний файл, KiB"],
      [(x["p"], x["files"], f(x["total_mib"], 2), f(x["avg_kib"], 1)) for x in S["small"]])

h("13. Экспериментальные выводы")
bm = {x["p"]: x for x in R["bench"]}
best = min(R["bench"], key=lambda x: x["median"])
concl = [
    f"Lazy evaluation: на стенде (local[*], Spark {R['stand']['PySpark / Spark']}) построение цепочки select→filter→withColumn над geolocation заняло {f(L['build_s'],4)} с и не создало Job ({L['jobs_after_transformations']}); count() создал Job и выполнялся {f(L['count_s'])} с — стоимость появляется только при Action.",
    f"Partition/Task: geolocation читается в {R['input_partitions']} партиций, первая стадия использует такое же число Tasks; число Tasks зависит от стадии (например, {G['Tasks по Stages']} у groupBy seller_id: 8 = spark.sql.shuffle.partitions).",
    f"Shuffle при groupBy: в плане есть Exchange hashpartitioning(seller_id, 8), Job разбился на Stages {G['Tasks по Stages']}, Shuffle Write = Shuffle Read = {f(G['Shuffle Write, KiB'],1)} КиБ.",
    f"repartition/coalesce: при уменьшении skewed16 до 4 партиций repartition(4) дал Exchange и выровнял размеры ({fi(rp['min'])}–{fi(rp['max'])} строк), coalesce(4) без полного shuffle сохранил перекос ({fi(co['min'])}–{fi(co['max'])}); на данных payment_type результат аналогичен ({fi(S['prof']['coalesce(4)']['min'])}–{fi(S['prof']['coalesce(4)']['max'])} у coalesce).",
    f"JOIN: при одинаковом результате ({fi(a['rows'])} строк) SMJ содержал {a['exchange']} Exchange и {a['sorts']} Sort и записал {f(a['Shuffle Write, KiB'],1)} КиБ shuffle, BHJ — BroadcastExchange без shuffle большой стороны ({f(b['Shuffle Write, KiB'],1)} КиБ); время {f(a['time'])} с против {f(b['time'])} с. На таблице 32 951 строк broadcast уместен; вывод не переносится на большие справочники.",
    f"cache/persist: с persist материализация заняла {f(c['mat'])} с, повторные Actions — {f(c['a1'])} и {f(c['a2'])} с против {f(nc['a1'])} и {f(nc['a2'])} без кэша. Суммарно кэш окупается только если повторных Actions достаточно, чтобы покрыть материализацию; для двух Actions выигрыш ограничен.",
    f"Зависимость от p: на N=2 000 000 минимальная медиана получена при p={best['p']} ({f(best['median'])} с); S(p) при p=32 составил {f(bm[32]['speedup'],2)}, т.е. при малом объёме вычислений рост числа партиций увеличил накладные расходы планирования Tasks, а не ускорил задачу. Монотонного ускорения нет.",
    f"Small files: одни и те же данные order_items при p=1/4/16/64 дали 1/4/16/64 part-файлов; средний размер файла упал с {f(R['small'][0]['avg_kib'],0)} до {f(R['small'][-1]['avg_kib'],0)} КиБ, а общий размер вырос с {f(R['small'][0]['total_mib'],2)} до {f(R['small'][-1]['total_mib'],2)} МиБ (хуже сжатие).",
    f"AQE: при включённом AQE план содержит AdaptiveSparkPlan, число Tasks стадий {R['aqe']['Tasks по Stages']} определяется рантайм-статистикой, поэтому основные опыты выполнялись с AQE=false.",
    f"Самостоятельная часть (orders ⋈ order_payments): SMJ записал {f(sa['Shuffle Write, KiB'],1)} КиБ shuffle против {f(sb['Shuffle Write, KiB'],1)} КиБ у BHJ при одинаковых {fi(sa['rows'])} строках; серия p=1/4/16/64 показала рост времени с {f(S['bench'][0]['median'])} до {f(S['bench'][-1]['median'])} с.",
]
for i, t in enumerate(concl, 1): p(f"{i}. {t}")

h("14. Ограничения измерений")
for t in [
    "Один компьютер, режим local[*]: все «исполнители» делят CPU, память и диск, поэтому сетевой стоимости shuffle нет, а абсолютное время несопоставимо со стендами других студентов.",
    "Малые объёмы (≤ 1 млн строк, 2 млн для benchmark) и короткие Jobs (доли секунды): различия сопоставимы с шумом JIT/GC и файлового кэша ОС; три повтора и медиана уменьшают, но не устраняют разброс. Heap JVM по умолчанию ~1 ГБ вызывал предупреждения MemoryManager при записи 64 партиций.",
    "Spark UI зафиксирован через REST API, а не скриншотами; на Windows для записи Parquet использован winutils/hadoop.dll версии 3.3.6.",
]: p("• " + t)

d.save("Отчет_Практика_4.docx")
print("ok")
